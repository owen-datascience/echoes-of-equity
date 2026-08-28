"""
Bug fix: the previous script's "Table 8 -> Table 6" substitution also matched
inside source-paper citations "[1, Table 8]". Restore them to [1, Table 8].
"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

SRC = "Owen-Wu-Yau-IEEE (1)-Reviewed by Duo Chen.docx"

def replace_paragraph_text(paragraph, new_text):
    for r in list(paragraph.runs):
        r._element.getparent().remove(r._element)
    paragraph.add_run(new_text)

d = Document(SRC)
n = 0
for p in d.paragraphs:
    txt = p.text
    if "[1, Table 6]" in txt:
        replace_paragraph_text(p, txt.replace("[1, Table 6]", "[1, Table 8]"))
        n += 1
        print(f"Reverted: {p.text[:150]}")

d.save(SRC)
print(f"\nRestored {n} citations. Saved: {SRC}")
