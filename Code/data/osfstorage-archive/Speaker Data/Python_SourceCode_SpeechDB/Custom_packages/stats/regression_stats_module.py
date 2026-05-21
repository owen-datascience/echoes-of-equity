"""
-----------------------------------------------------------------------------------------------------

The purpose of this module is to conduct classification and/or predictions using regression models
with the statistical package (statsmodels).

-----------------------------------------------------------------------------------------------------
"""

# Loading necessary package dependencies.
import pandas as pd
import numpy as np
import os.path
from docx import Document
import statsmodels.api as sm
from sklearn.model_selection import train_test_split
from sklearn.metrics import confusion_matrix, classification_report
from tabulate import tabulate
from Custom_packages.misc import visual_formatting_module, helpers_module


def calculate_logistic_regression(features: pd.DataFrame, target: pd.Series, group_name="") -> None:
    """
    Performs logistic regression (statsmodels package) on the provided dataset and exports the results to a DOCX file.

    Parameters:
        features (pd.DataFrame): Feature columns in the dataset.
        target (pd.Series): The dependent variable (target) for the regression.
        group_name (str, optional): A name for the demographic group, used in file naming and reports.

    Returns:
        None

    Example:
        calculate_logistic_regression(features, target, group_name="Younger")
    """
    try:
        helpers_module.validate_non_empty_and_exists(features, "features")
        helpers_module.validate_non_empty_and_exists(target,"target")

        model = sm.Logit(target, sm.add_constant(features))
        regression_results = model.fit()

        # Get a summary with the estimated information and diagnostics (e.g. goodness-of-fit, log-likelihood,
        # df residuals, coefficients (Beta), p-values, etc).
        summary_table_full = regression_results.summary2()
        summary_table_compact = regression_results.summary2().tables[1]

        # Calculate odds ratios for each acoustic feature based on the Coefficient estimates,
        # and add them to the summary table.
        summary_table_compact["Odds Ratios"] = np.exp(summary_table_compact["Coef."])

        # Export to DOCX
        export_LR_summary_table_to_MS_Word(summary_table_compact, group_name)

        formatted_summary = format_LR_summary_table(summary_table_compact)
        print(f"\n\nLogistic Regression for {group_name} speakers:\n")
        print("\nResults summary (full details):\n\n", summary_table_full)
        print("\nResults summary (compact version):\n\n", formatted_summary)

    except Exception as e:
        visual_formatting_module.print_coloured(f"\nError during statistical logistic regression\n: {e}", "red")


def export_LR_summary_table_to_MS_Word(summary_table_compact: pd.DataFrame, group_name: str) -> None:
    """
    Exports the logistic regression summary table to a DOCX file.

    Parameters:
        summary_table_compact (pd.DataFrame): The compact summary table to export.
        group_name (str): The group name used for the file name.

    Returns:
        None
    """
    current_dir = os.path.dirname(os.path.abspath(__file__))
    tables_export_folder_path = os.path.join(current_dir, "Exported_data_tables")

    if not os.path.exists(tables_export_folder_path):
        os.makedirs(tables_export_folder_path)

    # Round and prepare the table for export
    summary_table_compact = summary_table_compact.round(3)
    doc = Document()
    table_title = f"\n\nLogistic Regression summary table for '{group_name}' speakers:\n"
    doc.add_paragraph(table_title, style="Title")

    # Add table
    docx_table = doc.add_table(summary_table_compact.shape[0] + 1, summary_table_compact.shape[1] + 1)
    for i, col in enumerate(summary_table_compact.columns):
        docx_table.cell(0, i + 1).text = col

    for j, idx in enumerate(summary_table_compact.index):
        docx_table.cell(j + 1, 0).text = str(idx)
        for i in range(summary_table_compact.shape[1]):
            docx_table.cell(j + 1, i + 1).text = str(summary_table_compact.iloc[j, i])

    # Save the DOCX file
    file_path = os.path.join(tables_export_folder_path, f"LR_summary_table_{group_name}_speakers.docx")
    doc.save(file_path)

def format_LR_summary_table(summary_table_compact: pd.DataFrame) -> str:
    """
    Formats the logistic regression summary table for display, highlighting significant results.

    Parameters:
        summary_table_compact (pd.DataFrame): The logistic regression summary table.

    Returns:
        str: The formatted summary table.
    """
    summary_table_compact["P>|z|"] = summary_table_compact["P>|z|"].apply(visual_formatting_module.highlight_significant_results)
    return tabulate(summary_table_compact, headers="keys", tablefmt="plain")
