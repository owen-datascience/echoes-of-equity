"""
One-shot: apply items 1-7 from second review pass.
"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

SRC = "Owen-Wu-Yau-IEEE (1)-Reviewed by Duo Chen.docx"

def replace_paragraph_text(paragraph, new_text):
    for r in list(paragraph.runs):
        r._element.getparent().remove(r._element)
    paragraph.add_run(new_text)

d = Document(SRC)

changes = 0

for p in d.paragraphs:
    txt = p.text
    new = txt

    # Item 1: 5-seed caption fixes
    for old, new_str in [
        ("mean ± std over 5 seeds", "mean ± std over 50 seeds"),
        ("mean over 5 seeds", "mean over 50 seeds"),
        ("averaged over 5 seeds", "averaged over 50 seeds"),
        ("(5 seeds, same fixed split)", "(50 seeds, same fixed split)"),
    ]:
        if old in new:
            new = new.replace(old, new_str)

    # Item 2: RF acc 73.42% -> 73.05%
    if "73.42%" in new:
        new = new.replace("73.42%", "73.05%")

    # Item 3: Table 8 -> Table 6
    if "Table 8" in new:
        new = new.replace("Table 8", "Table 6")

    # Item 4: Section IV-G -> Section IV-F
    if "Section IV-G" in new:
        new = new.replace("Section IV-G", "Section IV-F")
    if "Sections IV-G" in new:
        new = new.replace("Sections IV-G", "Sections IV-F")
    if "Sections IV-B through IV-F" in new:
        new = new.replace("Sections IV-B through IV-F", "Sections IV-B through IV-E")

    # Item 4 (heading): "G. Ablation" -> "F. Ablation"
    if new.strip().startswith("G. Ablation"):
        new = new.replace("G. Ablation", "F. Ablation", 1)

    # Item 5: "Sections III-D, 3.5" typo
    if "Sections III-D, 3.5" in new:
        new = new.replace("Sections III-D, 3.5", "Sections III-D and III-E")

    # Item 7: drop "(Headline Result)" from IV-E header
    if new.strip().startswith("E. Variance and Stability Across Seeds"):
        new = new.replace(" (Headline Result)", "")

    if new != txt:
        replace_paragraph_text(p, new)
        changes += 1
        print(f"Changed [{p.style.name[:10]:10s}]: {new[:120]}")

# Item 6: Rewrite §IV overview paragraph
for i, p in enumerate(d.paragraphs):
    if p.text.strip().startswith("Section IV-A describes the shared experimental setup."):
        replace_paragraph_text(p, (
            "Section IV-A describes the shared experimental setup. Section IV-B "
            "presents overall classification performance (Table 2). Section IV-C "
            "decomposes accuracy by speaker ethnicity, the paper's primary "
            "fairness axis (Table 3, Fig. 6). Section IV-D summarizes AUC results "
            "by ethnicity (Table 4). Section IV-E quantifies the variance of every "
            "model's accuracy and fairness behavior across the 50 random "
            "initializations (Table 5, Fig. 7). Section IV-F reports an ablation "
            "on the DANN encoder architecture (Table 6). A per-age and per-sex "
            "breakdown supporting the sex-gap finding discussed in Section V "
            "appears as Fig. 8."
        ))
        changes += 1
        print("Item 6: §IV overview paragraph rewritten")
        break

d.save(SRC)
print(f"\nTotal changes: {changes}. Saved: {SRC}")
