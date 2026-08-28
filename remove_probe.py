"""
One-shot: remove all mentions of the frozen-encoder probe (Item 9 in review).
Also fix stale "five random seeds" -> "50 random seeds" in the §I-E paragraph.
"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

SRC = "Owen-Wu-Yau-IEEE (1)-Reviewed by Duo Chen.docx"

def replace_paragraph_text(paragraph, new_text):
    for r in list(paragraph.runs):
        r._element.getparent().remove(r._element)
    paragraph.add_run(new_text)

def delete_paragraph(paragraph):
    el = paragraph._element
    el.getparent().remove(el)

def find_paragraph_containing(d, substr):
    for p in d.paragraphs:
        if substr in p.text:
            return p
    return None

d = Document(SRC)

# 1. Abstract
p_abstract = find_paragraph_containing(d, "A logistic-regression probe trained on the frozen DANN encoder")
if p_abstract is not None:
    txt = p_abstract.text
    probe_sentence = (
        " A logistic-regression probe trained on the frozen DANN encoder still "
        "recovers ethnicity at ~53% against a 33.3% chance baseline, so "
        "ethnicity information is reduced in influence on the trust head but "
        "not erased."
    )
    new_txt = txt.replace(probe_sentence, "")
    replace_paragraph_text(p_abstract, new_txt)
    print("1. Abstract: probe sentence removed.")

# 2. §I-E: drop probe intro + fix seed count
p_ie = find_paragraph_containing(d, "We address both fairness and stability problems")
if p_ie is not None:
    replace_paragraph_text(p_ie, (
        "We address both fairness and stability problems using the TIS Corpus "
        "[20]. We evaluate five classifiers using a single 80/20 train-test "
        "split stratified jointly by trustworthy-intent label and ethnicity, "
        "ensuring that all models are evaluated on the same data partition. "
        "The models include the Random Forest and Logistic Regression baselines "
        "from the source study, a three-layer artificial neural network (ANN), "
        "a one-dimensional convolutional neural network (1D-CNN), and "
        "DANN-Trust, a domain-adversarial network incorporating a "
        "gradient-reversal layer [26]. Each model is retrained under 50 random "
        "seeds. Overall accuracy, the per-ethnicity fairness gap Δ_eth, and "
        "per-age and per-sex accuracy gaps are therefore evaluated as "
        "distributions rather than as single values, with pairwise "
        "Mann-Whitney U and Levene's tests supporting each claim. Together, "
        "these analyses move beyond the question of which model achieves the "
        "highest accuracy in a single experiment; they instead ask whether a "
        "model's demographic performance is balanced across ethnicity, age, "
        "and sex, and whether that balance persists when the model is "
        "retrained."
    ))
    print("2. §I-E: probe intro removed + seed count fixed (5 -> 50).")

# 3. §III-G probe methodology paragraph: delete
p_iiig = find_paragraph_containing(
    d, "For DANN-Trust, we apply a direct diagnostic on the adversarial mechanism"
)
if p_iiig is not None:
    delete_paragraph(p_iiig)
    print("3. §III-G probe methodology paragraph deleted.")

# 4. §IV-F probe results paragraph: delete
p_ivf = find_paragraph_containing(
    d, "To verify the adversarial mechanism directly, the frozen encoder output"
)
if p_ivf is not None:
    delete_paragraph(p_ivf)
    print("4. §IV-F probe results paragraph deleted.")

# 5. §V-B: rewrite without probe references
p_vb = find_paragraph_containing(d, "The frozen-encoder probe qualifies but does not overturn")
if p_vb is not None:
    replace_paragraph_text(p_vb, (
        "On this corpus and under the accuracy-gap metric used here, both CNN "
        "and DANN-Trust produce among the smallest per-ethnicity disparities "
        "observed while maintaining competitive overall accuracy. Neither model "
        "should be described as demographically invariant or universally fair "
        "in a stronger sense: the Δ metric measures only equalized per-group "
        "accuracy on the trust decision, not representation-level invariance, "
        "per-group error types, or equalized odds. The case for choosing CNN "
        "over DANN-Trust (or vice versa) on this corpus therefore rests on "
        "considerations outside the measured metrics — implementation "
        "complexity, interpretability, and behavior on future related datasets "
        "— since the two are statistically indistinguishable on every metric "
        "reported here. Whether the DANN-Trust encoder's internal "
        "representation is demographically invariant at the representation "
        "level (as opposed to at the decision level) is not evaluated here and "
        "remains an open question for future work."
    ))
    print("5. §V-B: rewritten without probe references.")

d.save(SRC)
print(f"\nSaved: {SRC}")
