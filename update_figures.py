"""
One-shot: swap the data-driven figures in the docx with the freshly-regenerated
50-seed PNGs, and fix the Fig. 6 caption "5 seeds" -> "50 seeds".

Static architecture diagrams (Fig 2, 3, 4, 5) are left untouched.

Mapping determined by walking the doc XML:
  rId13 -> image1.png -> Fig 1 (boxplots)            <- fig03_feature_boxplots.png
  rId18 -> image6.png -> Fig 6 (per-ethnicity acc)   <- fig09_acc_by_ethnicity.png
  rId19 -> image7.png -> Fig 7 (fairness pareto)     <- fig13_fairness_pareto.png
"""
from docx import Document

SRC = "Owen-Wu-Yau-IEEE (1) with Data Updated.docx"
FIG_DIR = "Paper/figures"

REPLACEMENTS = {
    "rId13": f"{FIG_DIR}/fig03_feature_boxplots.png",
    "rId18": f"{FIG_DIR}/fig09_acc_by_ethnicity.png",
    "rId19": f"{FIG_DIR}/fig13_fairness_pareto.png",
}

d = Document(SRC)
doc_part = d.part

for rid, new_path in REPLACEMENTS.items():
    if rid not in doc_part.rels:
        print(f"WARN: {rid} not in doc rels - skipping")
        continue
    rel = doc_part.rels[rid]
    target_part = rel.target_part
    with open(new_path, "rb") as f:
        new_bytes = f.read()
    target_part._blob = new_bytes
    print(f"Replaced {target_part.partname} ({rid}) with {new_path}  ({len(new_bytes):,} bytes)")

target_caption_idx = None
for i, p in enumerate(d.paragraphs):
    if p.text.strip().startswith("Fig. 6 Per-ethnicity"):
        target_caption_idx = i
        break

if target_caption_idx is not None:
    p = d.paragraphs[target_caption_idx]
    new_caption = (
        "Fig. 6 Per-ethnicity trust accuracy, mean +/- std over 50 seeds. "
        "Deep models (ANN/CNN/DANN) raise South Asian accuracy by ~8-10 pp "
        "over Random Forest, largely closing the baseline gap."
    )
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    p.add_run(new_caption)
    print(f"Updated Fig. 6 caption at paragraph [{target_caption_idx}].")
else:
    print("WARN: could not locate Fig. 6 caption.")

d.save(SRC)
print(f"\nWrote: {SRC}")
