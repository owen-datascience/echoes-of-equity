"""
One-shot: insert fig10_acc_by_age_sex.png into the docx as Fig. 8,
placed immediately before the "sex axis" paragraph in section V-A.
Adds a caption below the image and prefaces the sex-gap paragraph
with an in-text pointer.
"""
from copy import deepcopy
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph

SRC = "Owen-Wu-Yau-IEEE (1)-Reviewed by Duo Chen.docx"
NEW_IMG = "Paper/figures/fig10_acc_by_age_sex.png"

d = Document(SRC)

target_idx = None
for i, p in enumerate(d.paragraphs):
    if p.text.strip().startswith("A separate finding at 50 seeds concerns the sex axis"):
        target_idx = i
        break

if target_idx is None:
    raise SystemExit("Could not locate the sex-gap paragraph.")

sex_para = d.paragraphs[target_idx]

def insert_blank_paragraph_before(anchor_para):
    new_p = deepcopy(anchor_para._p)
    for child in list(new_p):
        new_p.remove(child)
    anchor_para._p.addprevious(new_p)
    return Paragraph(new_p, anchor_para._parent)

img_para = insert_blank_paragraph_before(sex_para)
img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = img_para.add_run()
run.add_picture(NEW_IMG, width=Inches(6.5))

caption_para = insert_blank_paragraph_before(sex_para)
caption_para.add_run(
    "Fig. 8  Per-age and per-sex trust accuracy, mean +/- std over 50 seeds. "
    "The male-speaker bar (right) narrows visibly under CNN and DANN-Trust "
    "relative to Random Forest and Logistic Regression, mirroring the "
    "significant sex-gap reductions reported in the text (Mann-Whitney U, "
    "p=0.003 for CNN, p=0.002 for DANN)."
)

POINTER = (
    "Figure 8 disaggregates overall accuracy by age group and by speaker sex. "
)
old_text = sex_para.text
for r in list(sex_para.runs):
    r._element.getparent().remove(r._element)
sex_para.add_run(POINTER + old_text)

d.save(SRC)
print(f"Inserted Fig. 8 image + caption before paragraph [{target_idx}].")
print(f"Prepended in-text pointer to the sex-gap paragraph.")
print(f"Wrote: {SRC}")
